"""Extractor seguro del Directorio de Empresas CCB.

Este script NO toca leads_prospeccion.xlsx ni envia mensajes. Genera un archivo separado
para revisar la calidad de los datos antes de integrarlos al flujo principal.
"""

from __future__ import annotations

import argparse
import html
import itertools
import json
import os
import re
import time
import warnings
from dataclasses import dataclass
from datetime import datetime
from typing import Iterable
from urllib.parse import urljoin, quote

warnings.filterwarnings("ignore", message="urllib3 v2 only supports OpenSSL")

CHECKPOINT_FILE = "ccb_checkpoint.json"


def load_checkpoint(path: str) -> int:
    """Devuelve la ultima pagina completada, o 0 si no hay checkpoint."""
    try:
        with open(path) as f:
            return int(json.load(f).get("last_page", 0))
    except (FileNotFoundError, ValueError, KeyError):
        return 0


def save_checkpoint(path: str, page_num: int) -> None:
    with open(path, "w") as f:
        json.dump({"last_page": page_num}, f)

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from core.config import COUNTRY_PREFIX, OUTPUT_FILE

BASE_URL = "https://negocios.ccb.org.co"
DIRECTORY_URL = f"{BASE_URL}/2-directorio-de-empresas"
OUTPUT_CCB = "leads_ccb_cuarentena.xlsx"
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0.0.0 Safari/537.36"
)

# Palabras que identifican empresas tech (proveedores de tecnología, no clientes)
TECH_EXCLUSION = [
    "software", "tecnolog", "developer", "desarroll web", "programac",
    "sistemas inform", "soluciones ti", "soluciones it", "it consulting",
    "startup", "saas", "app movil", "aplicac movil", "hosting", "cloud",
    "cibersegurid", "inteligencia artificial", "ia y", "machine learning",
    "agencia digital", "marketing digital", "desarrollo digital",
    "servicios digitales", "infraestructura ti", "infraestructura it",
    "network", "cyber", "soft", "digital", "technolog", "data", "datos",
]

# Heurística conservadora: no queremos quemar tiempo en empresas grandes o muy corporativas.
LARGE_COMPANY_EXCLUSION = [
    "carvajal", "postobon", "bancolombia", "davivienda", "grupo aval",
    "grupo bolivar", "bolivar", "sura", "nutresa", "argos", "colsubsidio",
    "compensar", "avianca", "falabella", "alkosto", "olimpica", "exito",
    "kenvue", "permoda", "casalimpia", "confipetrol", "superbid",
    "turner & townsend", "joyco", "multinacional", "holding", "corporacion",
    "corporativo", "grupo empresarial", "international", "internacional",
    " s a ", " s.a ", " s.a.", "sociedad anonima", "s.a.s. bic", " sas bic",
]

_EMAIL_SUBJECT = "\u00bfSu p\u00e1gina web les est\u00e1 trayendo clientes o se los est\u00e1 quitando?"
_EMAIL_BODY = (
    "Hola, equipo de {nombre}:\n\n"
    "Les escribo porque noto que muchas empresas invierten tiempo en redes sociales, "
    "pero pierden a los clientes cuando llegan a una p\u00e1gina web lenta, desactualizada "
    "o que no funciona bien.\n\n"
    "En SIO, nuestro fuerte principal es el desarrollo web: arreglamos p\u00e1ginas que no "
    "est\u00e1n dando resultados o las construimos desde cero con alto rendimiento. Nos "
    "aseguramos de que tengan una arquitectura t\u00e9cnica s\u00f3lida y lista para captar "
    "contactos.\n\n"
    "A partir de esa base impecable, nos encargamos del resto para que no tengan que "
    "lidiar con m\u00faltiples proveedores: renovamos su logo, manejamos sus redes, "
    "aplicamos SEO y automatizamos sus tareas repetitivas para conseguirles clientes.\n\n"
    "Tengo un par de ideas sobre c\u00f3mo podr\u00edamos optimizar su plataforma actual. "
    "\u00bfTendr\u00edan 10 minutos la pr\u00f3xima semana para una llamada r\u00e1pida y se las comparto?\n\n"
    "Saludos,\n\n"
    "Juan Aguilar | SIO"
)

_WA_MESSAGE = (
    "Hola {nombre}, mi nombre es Juan Aguilar de SIO.\n\n"
    "Encontré su empresa en {busqueda} y quería compartirles una idea para mejorar su presencia digital."
)


def clean_phone(raw: str) -> str:
    cleaned = re.sub(r"[\s()\-\.\+]", "", raw)
    if cleaned.startswith("57"):
        cleaned = cleaned[2:]
    return re.sub(r"\D", "", cleaned)


def build_whatsapp_link(phone_clean: str, nombre: str = "", busqueda: str = "") -> str:
    if not phone_clean:
        return ""
    msg = _WA_MESSAGE.format(nombre=nombre, busqueda=busqueda)
    return f"https://wa.me/{COUNTRY_PREFIX}{phone_clean}?text={quote(msg)}"

COLUMNS = [
    "Nombre",
    "Categoría",
    "Canal recomendado",
    "Dirección",
    "Teléfono",
    "WhatsApp Directo",
    "Link WhatsApp",
    "Email",
    "Web",
    "NIT",
    "Descripción",
    "Búsqueda",
    "Fuente",
    "URL CCB",
    "Estado Dato",
    "Contactar",
]

PREMIUM_HINTS = [
    "abogado", "jurid", "legal", "consult", "asesor", "financier", "contador",
    "revisor", "auditor", "seguros", "inmobili", "construct", "arquitect",
    "ingenier", "tecnolog", "software", "infraestructura", "hotel", "salud",
]


@dataclass
class CompanyLink:
    name: str
    url: str


def _strip_tags(value: str) -> str:
    value = re.sub(r"<script\b.*?</script>", " ", value, flags=re.I | re.S)
    value = re.sub(r"<style\b.*?</style>", " ", value, flags=re.I | re.S)
    value = re.sub(r"<[^>]+>", " ", value)
    value = html.unescape(value)
    return re.sub(r"\s+", " ", value).strip()


def _session() -> requests.Session:
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT, "Accept-Language": "es-CO,es;q=0.9"})
    return session


def fetch(session: requests.Session, url: str) -> str:
    response = session.get(url, timeout=25)
    response.raise_for_status()
    return response.text


def load_existing_phones(filename: str = OUTPUT_FILE) -> set[str]:
    phones: set[str] = set()
    if not os.path.exists(filename):
        return phones
    try:
        wb = load_workbook(filename, read_only=True)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        if "WhatsApp Directo" not in headers:
            return phones
        phone_idx = headers.index("WhatsApp Directo")
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw = (row[phone_idx] or "") if phone_idx < len(row) else ""
            phone = str(raw).replace("+", "").strip()
            if phone:
                phones.add(phone)
        wb.close()
    except Exception as exc:
        print(f"WARN: no se pudieron cargar teléfonos existentes: {exc}")
    return phones


def parse_company_links(page_html: str) -> list[CompanyLink]:
    pattern = re.compile(
        r'<a[^>]+href="(?P<href>[^"]*/directorio-de-empresas/[^"]+\.html)"[^>]*>(?P<name>.*?)</a>',
        re.I | re.S,
    )
    links: list[CompanyLink] = []
    seen: set[str] = set()
    for match in pattern.finditer(page_html):
        url = urljoin(BASE_URL, html.unescape(match.group("href")))
        name = _strip_tags(match.group("name"))
        key = url.lower()
        if not name or key in seen:
            continue
        seen.add(key)
        links.append(CompanyLink(name=name, url=url))
    return links


def extract_emails(raw_html: str) -> list[str]:
    emails: list[str] = []
    seen: set[str] = set()
    mailto_pattern = re.compile(r"mailto:([A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})", flags=re.I)
    generic_pattern = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", flags=re.I)

    for pattern in (mailto_pattern, generic_pattern):
        for match in pattern.findall(raw_html):
            email = html.unescape(match).strip().lower()
            if email.endswith("@ccb.org.co") or email in seen:
                continue
            seen.add(email)
            emails.append(email)
    return emails


def extract_websites(raw_html: str, own_url: str) -> list[str]:
    websites: set[str] = set()
    for href in re.findall(r'href=["\']([^"\']+)["\']', raw_html, flags=re.I):
        href = html.unescape(href).strip()
        if not href.startswith(("http://", "https://")):
            continue
        if "negocios.ccb.org.co" in href or "ccb.org.co" in href or "lineaeticaccb.org" in href or href == own_url:
            continue
        if any(domain in href for domain in ("facebook.com", "instagram.com", "twitter.com", "linkedin.com", "youtube.com", "tiktok.com")):
            continue
        websites.add(href)
    return sorted(websites)


def extract_phone(raw_html: str) -> tuple[str, str]:
    candidates: list[str] = []
    candidates.extend(re.findall(r"(?:wa\.me/|phone=|whatsapp[^0-9]{0,30})(57?3\d{9})", raw_html, flags=re.I))
    candidates.extend(re.findall(r"(?:tel:|Tel[eé]fono:?\s*)(\+?57?\s*3[\d\s().-]{9,})", raw_html, flags=re.I))
    candidates.extend(re.findall(r"\+?57?\s*3\d[\d\s().-]{8,}", raw_html))

    for candidate in candidates:
        clean = clean_phone(candidate)
        if clean.startswith("57"):
            clean = clean[2:]
        if clean.startswith("3") and len(clean) == 10:
            full = f"{COUNTRY_PREFIX}{clean}"
            return clean, f"+{full}"
    return "", ""


def parse_detail(company: CompanyLink, detail_html: str) -> dict:
    title = company.name
    title_match = re.search(r"<h1[^>]*>(.*?)</h1>", detail_html, flags=re.I | re.S)
    if title_match:
        title = _strip_tags(title_match.group(1)) or title

    text = _strip_tags(detail_html)
    nit_match = re.search(r"NIT\s*([0-9.-]+)", text, flags=re.I)
    nit = nit_match.group(1).strip() if nit_match else ""

    description = ""
    title_pos = text.lower().find(title.lower())
    info_pos = text.lower().find("whatsapp", title_pos if title_pos >= 0 else 0)
    if title_pos >= 0 and info_pos > title_pos:
        description = text[title_pos + len(title):info_pos].strip(" -")[:800]

    local_phone, whatsapp = extract_phone(detail_html)
    emails = extract_emails(detail_html)
    websites = extract_websites(detail_html, company.url)

    searchable = f"{title} {description}".lower()
    premium = any(hint in searchable for hint in PREMIUM_HINTS)
    recommended_channel = "WhatsApp" if whatsapp else "Correo" if emails else "Revisar"
    status = "con telefono" if whatsapp else "con correo publico" if emails else "sin contacto visible"
    if "para poder continuar" in text.lower() and not whatsapp and not emails:
        status = "contacto protegido por login"

    email_addr = emails[0] if emails else ""
    body = _EMAIL_BODY.format(nombre=title.title())
    mailto = (
        f"mailto:{email_addr}?subject={quote(_EMAIL_SUBJECT)}&body={quote(body)}"
        if email_addr else ""
    )

    return {
        "Nombre": title,
        "Categoría": "CCB Premium" if premium else "CCB Directorio",
        "Canal recomendado": recommended_channel,
        "Dirección": "",
        "Teléfono": local_phone,
        "WhatsApp Directo": whatsapp,
        "Link WhatsApp": build_whatsapp_link(local_phone, title, "CCB Directorio de Empresas") if whatsapp else "",
        "Email": emails[0] if emails else "",
        "Web": websites[0] if websites else "",
        "NIT": nit,
        "Descripción": description,
        "Búsqueda": "CCB Directorio de Empresas",
        "Fuente": "CCB",
        "URL CCB": company.url,
        "Estado Dato": status,
        "Contactar": mailto,
    }


def export_rows(rows: list[dict], filename: str = OUTPUT_CCB) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "CCB cuarentena"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="7B1E3A", end_color="7B1E3A", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin")
    )

    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(COLUMNS, start=1):
            value = row.get(col_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            if col_name == "Contactar" and value:
                cell.value = "\u2709 Contactar"
                cell.hyperlink = value
                cell.font = Font(color="0B6E2E", bold=True, underline="single")
                continue
            if col_name in ("Link WhatsApp", "URL CCB", "Web") and value:
                cell.hyperlink = value
                cell.font = Font(color="1155CC", underline="single")
            if col_name == "Email" and value:
                cell.hyperlink = f"mailto:{value}"
                cell.font = Font(color="1155CC", underline="single")

    for col_idx, col_name in enumerate(COLUMNS, start=1):
        max_len = len(col_name)
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 4, 60)
    ws.auto_filter.ref = ws.dimensions

    tmp = filename + ".tmp"
    wb.save(tmp)
    os.replace(tmp, filename)


def scrape_pages(pages: Iterable[int], limit: int, delay: float, include_without_phone: bool, checkpoint_file: str = "") -> list[dict]:
    session = _session()
    existing_phones = load_existing_phones()
    rows: list[dict] = []
    seen_urls: set[str] = set()
    seen_phones: set[str] = set(existing_phones)

    try:
        for page_num in pages:
            url = DIRECTORY_URL if page_num == 1 else f"{DIRECTORY_URL}?page={page_num}"
            print(f"Pagina {page_num}: {url}")
            listing_html = fetch(session, url)
            companies = parse_company_links(listing_html)
            print(f"  Empresas encontradas: {len(companies)}")
            for company in companies:
                if limit and len(rows) >= limit:
                    return rows
                if company.url.lower() in seen_urls:
                    continue
                seen_urls.add(company.url.lower())
                try:
                    detail_html = fetch(session, company.url)
                    row = parse_detail(company, detail_html)
                    # Filtrar empresas tech por nombre; la descripción de CCB trae texto común de la página.
                    searchable = row["Nombre"].lower()
                    if any(kw in searchable for kw in TECH_EXCLUSION):
                        print(f"  TECH excluida: {row['Nombre']}")
                        continue
                    company_name_key = f" {searchable} "
                    if any(kw in company_name_key for kw in LARGE_COMPANY_EXCLUSION):
                        print(f"  GRANDE excluida: {row['Nombre']}")
                        continue
                    if row.get("Web"):
                        print(f"  Con web excluida: {row['Nombre']}")
                        continue
                    phone_key = row["WhatsApp Directo"].replace("+", "")
                    if phone_key and phone_key in seen_phones:
                        print(f"  DUP teléfono: {row['Nombre']}")
                        continue
                    if not phone_key and not row.get("Email") and not include_without_phone:
                        print(f"  Sin teléfono visible: {row['Nombre']}")
                        continue
                    if phone_key:
                        seen_phones.add(phone_key)
                    rows.append(row)
                    print(f"  + {row['Nombre']} [{row['Estado Dato']}]")
                except Exception as exc:
                    print(f"  ERROR {company.name}: {exc}")
                time.sleep(delay)
            if checkpoint_file:
                save_checkpoint(checkpoint_file, page_num)
    except KeyboardInterrupt:
        print(f"\n[Ctrl+C] Detenido en pagina {page_num}. Exportando {len(rows)} leads...")
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="Extractor seguro CCB a Excel separado")
    parser.add_argument("--start-page", type=int, default=None, help="página inicial (ignora checkpoint si se da)")
    parser.add_argument("--limit", type=int, default=0, help="máximo registros (0 = sin limite, detener con Ctrl+C)")
    parser.add_argument("--delay", type=float, default=1.5, help="pausa entre perfiles")
    parser.add_argument("--output", default=OUTPUT_CCB, help="Excel de salida separado")
    parser.add_argument("--include-without-phone", action="store_true", help="guardar perfiles sin teléfono visible")
    parser.add_argument("--reset", action="store_true", help="borrar checkpoint y empezar desde pagina 1")
    args = parser.parse_args()

    if args.reset:
        if os.path.exists(CHECKPOINT_FILE):
            os.remove(CHECKPOINT_FILE)
            print("Checkpoint borrado. Empezando desde pagina 1.")
        start_page = 1
    elif args.start_page is not None:
        start_page = args.start_page
    else:
        last = load_checkpoint(CHECKPOINT_FILE)
        start_page = last + 1
        if last:
            print(f"Retomando desde pagina {start_page} (ultima completada: {last})")

    print("Corriendo. Presiona Ctrl+C en cualquier momento para detener y guardar el Excel.")
    pages = itertools.count(start_page)
    rows = scrape_pages(pages, args.limit, args.delay, args.include_without_phone, CHECKPOINT_FILE)
    export_rows(rows, args.output)
    with_phone = sum(1 for row in rows if row.get("WhatsApp Directo"))
    with_email = sum(1 for row in rows if row.get("Email"))
    print(f"Guardado {args.output}: {len(rows)} registros, {with_phone} con teléfono/WhatsApp visible, {with_email} con correo")


if __name__ == "__main__":
    main()
